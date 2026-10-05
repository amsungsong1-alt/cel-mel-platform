"""Generates the SAWA Technical Narrative Report (.docx) from live app data."""
from __future__ import annotations
import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from database.db import run_query

PH = "[PLACEHOLDER — data not yet available]"


def _shd(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _hrow(row, *headers, bg="4A5568"):
    for i, h in enumerate(headers):
        if i >= len(row.cells):
            break
        c = row.cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shd(c, bg)


def _v(val, fb="—") -> str:
    if val is None:
        return fb
    s = str(val).strip()
    return s if s else fb


def _qtr_period(q: int, fy: int) -> str:
    return {
        1: f"July–September {fy}",
        2: f"October–December {fy}",
        3: f"January–March {fy + 1}",
        4: f"April–June {fy + 1}",
    }.get(q, "")


def build_report(project_id: int, fiscal_year: int, quarter: int) -> bytes:
    doc = Document()

    # A4 page setup
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    period = _qtr_period(quarter, fiscal_year)
    fy_str = f"FY{fiscal_year}/{str(fiscal_year + 1)[-2:]}"
    q_col = f"actual_q{quarter}"

    # ── Data fetch ────────────────────────────────────────────────────────────
    pr = run_query(
        "SELECT * FROM participant_register WHERE project_id=:p AND fiscal_year=:fy AND quarter=:q",
        {"p": project_id, "fy": fiscal_year, "q": quarter},
    )
    pr_all = run_query(
        "SELECT * FROM participant_register WHERE project_id=:p AND fiscal_year=:fy",
        {"p": project_id, "fy": fiscal_year},
    )
    mel = run_query(
        """SELECT ma.*, p.name AS partner_name FROM mel_activities ma
           LEFT JOIN partners p ON p.partner_id = ma.partner_id
           WHERE ma.project_id=:p AND ma.fiscal_year=:fy AND ma.quarter=:q
           ORDER BY ma.activity_date""",
        {"p": project_id, "fy": fiscal_year, "q": quarter},
    )
    lrn = run_query(
        "SELECT * FROM learnings WHERE project_id=:p AND fiscal_year=:fy AND quarter=:q ORDER BY learning_type, id",
        {"p": project_id, "fy": fiscal_year, "q": quarter},
    )
    qra = run_query(
        "SELECT * FROM qr_actuals WHERE project_id=:p AND fiscal_year=:fy ORDER BY submitting_org, indicator_code",
        {"p": project_id, "fy": fiscal_year},
    )
    tools = run_query(
        "SELECT instrument_name, stakeholder, frequency, responsible_party, rationale FROM data_collection_plan WHERE project_id=:p ORDER BY instrument_name",
        {"p": project_id},
    )
    rda = run_query(
        """SELECT lr.indicator_code, rda.actual_q1, rda.actual_q2, rda.actual_q3, rda.actual_q4, rda.target_value
           FROM raw_data_analysis rda
           JOIN logframe_rows lr ON lr.id = rda.logframe_row_id
           WHERE rda.project_id=:p AND (rda.reporting_year=:fy OR rda.reporting_year IS NULL)
             AND rda.partner_id IS NULL""",
        {"p": project_id, "fy": fiscal_year},
    )

    rda_map = {r["indicator_code"]: r for r in rda}
    qra_cel = {r["indicator_code"]: r for r in qra if r.get("submitting_org") == "CEL"}

    def _act(code: str) -> str:
        r = qra_cel.get(code)
        if r:
            v = r.get(q_col)
            if v and str(v).strip():
                return str(v)
        r2 = rda_map.get(code)
        if r2:
            v = r2.get(q_col)
            if v and str(v).strip():
                return str(v)
        return "—"

    def cnt(rows, **kw):
        r = rows
        for k, v in kw.items():
            r = [x for x in r if x.get(k) == v]
        return len(r)

    n = len(pr)
    n_f = cnt(pr, sex="Female")
    n_pwd = cnt(pr, is_pwd="Y")
    n_pwd_f = cnt([r for r in pr if r.get("is_pwd") == "Y"], sex="Female")
    n_rural = cnt(pr, location_type="Rural")
    n_urban = cnt(pr, location_type="Urban")
    n_peri = cnt(pr, location_type="Peri-urban")
    n_wan = cnt(pr, wan_member="Y")
    n_refugee = cnt(pr, refugee_displaced="Y")
    n_refugee_f = cnt([r for r in pr if r.get("refugee_displaced") == "Y"], sex="Female")
    n_regions = len({r.get("region") for r in pr_all if r.get("region")})
    n_districts = len({r.get("district") for r in pr_all if r.get("district")})
    n_communities = len({r.get("community") for r in pr_all if r.get("community")})

    weo_val = _act("PI.1") if _act("PI.1") != "—" else (str(n) if n > 0 else "—")
    yiw_val = _act("PIII.R1")
    pwd_yiw = _act("PII.R5")

    def sv(val: int) -> str:
        return str(val) if val > 0 else "—"

    # ═══════════════════════════════════════════════════════════════════════════
    # COVER
    # ═══════════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SAWA PROGRAMME")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Technical Narrative Report")
    r.bold = True
    r.font.size = Pt(16)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Q{quarter} — {period}")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0xF5, 0xA8, 0x20)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Centre for Entrepreneurship and Leadership (CEL)\n"
        "Sustainable Aquaculture for Wealth and Achievement (SAWA) Programme\n"
        f"Reporting Period: {period}  ·  {fy_str}\n"
        f"Generated: {datetime.utcnow().strftime('%d %B %Y')}"
    )

    doc.add_page_break()

    # ═══════════════════════════════════════════════════════════════════════════
    # SUMMARY
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("SUMMARY", level=1)
    p = doc.add_paragraph(
        "In ¼ to ½ a page, summarise the most important achievements during the reporting period "
        "and key elements for the Consolidated Annual Report."
    )
    p.italic = True
    doc.add_paragraph(PH)

    doc.add_heading("Key Highlights for the Quarter", level=2)

    _SECTION_HDRS = {
        "AGGREGATE", "OPERATIONAL REGIONS, DISTRICTS & COMMUNITIES",
        "YOUTH IN WORK", "WOMEN IN WORK", "EMPLOYMENT CATEGORY", "EMPLOYMENT TYPE",
        "REFUGEE AND/OR DISPLACED PERSONS", "PERSONS WITH DISABILITIES",
        "WOMEN IN AQUACULTURE NETWORK (WAN)", "ENTERPRISES SUPPORTED",
    }

    summary_data = [
        ("AGGREGATE", None),
        ("Work Enabling Outreach (WEO) — Total Youth", weo_val),
        ("Total Youth women 18–35 (accessing training)", sv(n_f)),
        ("Total Youth women 18–35 (accessing starter pack)", "—"),
        ("Total Rural", sv(n_rural)),
        ("Total Urban", sv(n_urban)),
        ("Total Peri-Urban", sv(n_peri)),
        ("OPERATIONAL REGIONS, DISTRICTS & COMMUNITIES", None),
        ("Total Regions", sv(n_regions)),
        ("Total Districts", sv(n_districts)),
        ("Total Communities", sv(n_communities)),
        ("YOUTH IN WORK", None),
        ("Total Youth in Work", yiw_val),
        ("Total Youth in Work — Primary", "—"),
        ("Total Youth in Work — Secondary", "—"),
        ("WOMEN IN WORK", None),
        ("Total Women Youth in Work", "—"),
        ("Total Women Youth in Work — Primary", "—"),
        ("Total Women Youth in Work — Secondary", "—"),
        ("EMPLOYMENT CATEGORY", None),
        ("Total Self Employment", "—"),
        ("Total Wage Employment", "—"),
        ("EMPLOYMENT TYPE", None),
        ("Total New Employment", "—"),
        ("Total Improved Employment", "—"),
        ("Total Additional Employment", "—"),
        ("REFUGEE AND/OR DISPLACED PERSONS", None),
        ("Total Refugee and/or Displaced Persons", sv(n_refugee)),
        ("Total Women Refugee and/or Displaced Persons", sv(n_refugee_f)),
        ("Total Refugee and/or Displaced Persons in Work", "—"),
        ("PERSONS WITH DISABILITIES", None),
        ("Total Youth with Disabilities (registered)", sv(n_pwd)),
        ("Total Young Women with Disabilities", sv(n_pwd_f)),
        ("Total Youth with Disabilities in Work", pwd_yiw),
        ("WOMEN IN AQUACULTURE NETWORK (WAN)", None),
        ("Number of WAN members (this quarter)", sv(n_wan)),
        ("Number of WAN formed — Total", "—"),
        ("ENTERPRISES SUPPORTED", None),
        ("Number of Enterprises Supported — Total", "—"),
        ("Number of Enterprises Supported — Women-Led", "—"),
        ("Number of Enterprises Supported — Youth-Led", "—"),
    ]

    t = doc.add_table(rows=1 + len(summary_data), cols=2)
    t.style = "Table Grid"
    _hrow(t.rows[0], "Measure / Indicator", "Count / Actuals", bg="F5A820")
    # Fix gold header text to dark
    for cell in t.rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

    for i, (label, val) in enumerate(summary_data, 1):
        row = t.rows[i]
        if label in _SECTION_HDRS:
            merged = row.cells[0].merge(row.cells[1])
            merged.text = ""
            run = merged.paragraphs[0].add_run(label)
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shd(merged, "4A5568")
        else:
            row.cells[0].text = label
            row.cells[1].text = val or "—"
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════════════════════════════
    # §1 PROGRAM DELIVERY
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("1.  PROGRAM DELIVERY MODELS OR APPROACHES", level=1)
    for num, title in [
        ("1.1.", "Models and Approaches"),
        ("1.2.", "Participant Attraction Strategy"),
        ("1.3.", "Pathways to Work"),
        ("1.4.", "Political Environment and Socioeconomic Developments"),
    ]:
        doc.add_heading(f"{num} {title}", level=2)
        doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # §2 KEY ACTIVITIES
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("2.  KEY ACTIVITIES AND OUTPUTS", level=1)

    # ── §2.1 Gender & Inclusion ───────────────────────────────────────────────
    doc.add_heading("2.1.  Young Women-Centred Capacity Building and Inclusion", level=2)
    doc.add_heading("2.1.1.  Gender", level=3)

    # Gender activity table
    g_acts = [r for r in mel if r.get("activity_type") in ("Learning session", "FGD")]
    n_gat = max(len(g_acts), 3)
    gat = doc.add_table(rows=1 + n_gat, cols=6)
    gat.style = "Table Grid"
    _hrow(gat.rows[0], "Gender Specific Activity", "Topic", "Objective(s)",
          "Achieved Outcome(s)", "Total Participants", "Female Participants")
    for i in range(1, 1 + n_gat):
        if i - 1 < len(g_acts):
            a = g_acts[i - 1]
            vals = [_v(a.get("activity_type")), "—",
                    _v(a.get("objective"))[:200], _v(a.get("key_findings"))[:200],
                    _v(a.get("participants")), "—"]
        else:
            vals = ["—"] * 6
        for j, v in enumerate(vals):
            gat.rows[i].cells[j].text = v
            for para in gat.rows[i].cells[j].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)

    doc.add_paragraph()
    p = doc.add_paragraph("Key Achievements on Gender Parity", style="Intense Quote")
    if n_f > 0:
        doc.add_paragraph(
            f"During {period}, {n_f} young women were registered as participants "
            f"out of {n} total ({n_f / max(n, 1) * 100:.0f}% female). " + PH
        )
    else:
        doc.add_paragraph(PH)

    p = doc.add_paragraph("Gender and Safeguarding Challenges and Barriers", style="Intense Quote")
    doc.add_paragraph(PH)
    p = doc.add_paragraph("Most Significant Incidents (Gender & Inclusion)", style="Intense Quote")
    p2 = doc.add_paragraph("Report at least 3 incidents. Include at least one PWD-related incident.")
    p2.italic = True
    doc.add_paragraph(PH)

    # §2.1.2 Safeguarding checklist
    doc.add_heading("2.1.2.  SAWA Safeguarding Checklist", level=3)
    p = doc.add_paragraph("Prepared by Gender & Safeguarding Lead.")
    p.italic = True

    sg_issues = [
        "Sexual harassment", "Sexual abuse and exploitation",
        "Sexist jokes and unwanted flirting", "Physical abuse", "Financial abuse",
        "Emotional and psychological abuse", "Verbal abuse", "Gender discrimination",
        "Ethnic/religious discrimination", "Act of omission and neglect",
        "Child labour", "Child abuse and exploitation", "Child protection issues",
        "Forced labour", "Modern day slavery",
    ]
    sgt = doc.add_table(rows=1 + len(sg_issues), cols=5)
    sgt.style = "Table Grid"
    _hrow(sgt.rows[0], "Key Safeguarding Issue", "Potential Risk",
          "Severity of Risk", "Risk Mitigation Measures", "Mitigation Effectiveness")
    for i, issue in enumerate(sg_issues, 1):
        sgt.rows[i].cells[0].text = issue
        for j in range(1, 5):
            sgt.rows[i].cells[j].text = "—"
        for cell in sgt.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()

    # §2.1.3 Youth Voice
    doc.add_heading("2.1.3.  Youth Voice and Agency", level=3)
    fgds = [r for r in mel if r.get("activity_type") == "FGD"]
    if fgds:
        doc.add_paragraph(f"{len(fgds)} FGD(s) conducted this quarter:")
        for r in fgds:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(f"{r.get('activity_date', '—')} · {r.get('location', '—')}:").bold = True
            bp.add_run(f" {_v(r.get('key_findings'))}")
    else:
        doc.add_paragraph(PH)

    # §2.1.4 PWD Inclusion
    doc.add_heading("2.1.4.  PWD Inclusion", level=3)
    if n_pwd > 0:
        doc.add_paragraph(
            f"{n_pwd} person(s) with disabilities registered this quarter "
            f"({n_pwd_f} female, {n_pwd - n_pwd_f} male — "
            f"{n_pwd / max(n, 1) * 100:.1f}% of total participants). "
            f"D&F work access (PII.R5 Q{quarter} actual): {pwd_yiw}."
        )
    else:
        doc.add_paragraph(PH)
    p = doc.add_paragraph("Barriers to PWDs participation:", style="Intense Quote")
    doc.add_paragraph(PH)
    p = doc.add_paragraph("Measures taken to increase participation of PWDs:", style="Intense Quote")
    doc.add_paragraph(PH)

    # ── §2.2 Production ───────────────────────────────────────────────────────
    doc.add_heading("2.2.  Production Expansion and Productivity Enhancement", level=2)
    p = doc.add_paragraph("Prepared by implementing partner. MEAL provides verification data from field visits and Kobo submissions.")
    p.italic = True

    cb_labels = [
        "Tilapia — Maize (Feed)", "Tilapia — Soybean (Feed)", "Tilapia — Hatchery",
        "Tilapia — Fingerlings", "Tilapia — Grow-out",
        "Catfish — Maize (Feed)", "Catfish — Soybean (Feed)", "Catfish — Hatchery",
        "Catfish — Fingerlings", "Catfish — Grow-out",
    ]
    cbt = doc.add_table(rows=1 + len(cb_labels), cols=6)
    cbt.style = "Table Grid"
    _hrow(cbt.rows[0], "#", "Description of Training Delivery", "Facilitator",
          "Female Participants", "Male Participants", "Key Outputs")
    for i, label in enumerate(cb_labels, 1):
        cbt.rows[i].cells[0].text = str(i)
        cbt.rows[i].cells[1].text = label
        for j in range(2, 6):
            cbt.rows[i].cells[j].text = "—"
        for cell in cbt.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()

    doc.add_heading("2.2.1.  Tilapia", level=3)
    doc.add_paragraph(PH)
    doc.add_heading("2.2.2.  Catfish", level=3)
    doc.add_paragraph(PH)

    # ── §2.3 Value Addition ───────────────────────────────────────────────────
    doc.add_heading("2.3.  Value Addition, Market Systems and Entrepreneurship", level=2)
    doc.add_heading("2.3.1.  Tilapia", level=3)
    doc.add_paragraph(PH)
    doc.add_heading("2.3.2.  Catfish", level=3)
    doc.add_paragraph(PH)

    # ── §2.4 Ecosystem ────────────────────────────────────────────────────────
    doc.add_heading("2.4.  Ecosystem Strengthening", level=2)
    doc.add_heading("2.4.1.  Women in Aquaculture Enterprises (WAN)", level=3)
    if n_wan > 0:
        doc.add_paragraph(f"{n_wan} WAN member(s) registered this quarter.")
    else:
        doc.add_paragraph(PH)

    # ── §2.5 Crosscutting ─────────────────────────────────────────────────────
    doc.add_heading("2.5.  Crosscutting", level=2)
    doc.add_heading("2.5.1.  E-SAWA / Digital Technology", level=3)
    esawa = [r for r in pr if r.get("intervention_type") == "E-SAWA"]
    if esawa:
        doc.add_paragraph(f"{len(esawa)} participant(s) enrolled in E-SAWA this quarter.")
    else:
        doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # §3 PROGRAM MEL — MEAL OWNS
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("3.  PROGRAM MEL (MONITORING, EVALUATION AND LEARNING)", level=1)

    # §3.1 MEL Activities
    doc.add_heading("3.1.  Quarterly MEL Activities", level=2)
    if mel:
        melt = doc.add_table(rows=1 + len(mel), cols=5)
        melt.style = "Table Grid"
        _hrow(melt.rows[0], "S/N", "MEL Activity", "Location",
              "Objective", "Key Findings & Recommendations")
        for i, a in enumerate(mel, 1):
            melt.rows[i].cells[0].text = str(i)
            melt.rows[i].cells[1].text = f"{_v(a.get('activity_date'))} — {_v(a.get('activity_type'))}"
            melt.rows[i].cells[2].text = _v(a.get("location"))
            melt.rows[i].cells[3].text = _v(a.get("objective"))[:300]
            melt.rows[i].cells[4].text = _v(a.get("key_findings"))[:500]
            for cell in melt.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8)
    else:
        doc.add_paragraph(PH)

    # §3.2
    doc.add_heading("3.2.  Work with Anchor Partners", level=2)
    doc.add_paragraph(PH)

    # §3.3 Training Outcomes
    doc.add_heading("3.3.  Training and Capacity Development Outcomes Assessment", level=2)
    ppt = [r for r in mel if r.get("activity_type") == "Pre/post test"]
    if ppt:
        doc.add_paragraph(f"{len(ppt)} pre/post test activity/ies conducted:")
        for r in ppt:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(f"{_v(r.get('activity_date'))} · {_v(r.get('location'))}:").bold = True
            bp.add_run(f" {_v(r.get('key_findings'))}")
    else:
        doc.add_paragraph(PH)

    # §3.4 Monitoring Tools
    doc.add_heading("3.4.  Monitoring Tools and Methodology", level=2)
    if tools:
        tt = doc.add_table(rows=1 + len(tools), cols=5)
        tt.style = "Table Grid"
        _hrow(tt.rows[0], "Instrument", "Data Source / Stakeholder",
              "Frequency", "Responsible Party", "Rationale")
        for i, tool in enumerate(tools, 1):
            tt.rows[i].cells[0].text = _v(tool.get("instrument_name"))
            tt.rows[i].cells[1].text = _v(tool.get("stakeholder"))
            tt.rows[i].cells[2].text = _v(tool.get("frequency"))
            tt.rows[i].cells[3].text = _v(tool.get("responsible_party"))
            tt.rows[i].cells[4].text = _v(tool.get("rationale"))[:300]
            for cell in tt.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8)
    else:
        doc.add_paragraph(PH)

    # §3.5 Data Storage
    doc.add_heading("3.5.  Data Storage and Accessibility", level=2)
    doc.add_paragraph(
        "Raw data is stored in KoboToolbox (EU server). "
        "Exports, processed data and reports are filed in the CEL SharePoint/OneDrive SAWA folder."
    )
    doc.add_paragraph(f"SharePoint/OneDrive folder URL: {PH}")

    # ═══════════════════════════════════════════════════════════════════════════
    # §4 COMMUNICATION
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("4.  COMMUNICATION", level=1)
    p = doc.add_paragraph("Prepared by Communications team. MEAL verifies that story numbers match the data.")
    p.italic = True
    for num, title in [
        ("4.1.", "Introduction"), ("4.2.", "Highlights & Impact"),
        ("4.3.", "Participant's Impact Story"), ("4.4.", "Branding & Visibility Compliance"),
        ("4.5.", "Challenges & Lessons Learned"), ("4.6.", "Next Quarter Priorities"),
        ("4.7.", "Links to Stories & Social Media Updates"),
    ]:
        doc.add_heading(f"{num} {title}", level=2)
        doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # §5 RISK
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("5.  RISK", level=1)

    _risks = [
        ("1. Program level risk — unintended consequences, conflict of interest, market distortion",
         "Low", "Medium", PH, PH, PH),
        ("2. Program level risk — limited external evidence the program will achieve desired outcomes",
         "Medium", "High",
         f"MEAL conducting systematic data collection via KoboToolbox; participant register maintained; "
         f"quarterly reconciliation against Y1 targets (Module L). PI.1 Q{quarter} actual: {weo_val}.",
         "Medium",
         "MEAL data collection on track. Evidence base strengthening as programme scales."),
        ("3. Safeguarding risk", "Low", "High", PH, PH, PH),
        ("4. Risk to active participation of young women", "Medium", "High", PH, PH, PH),
        ("5. Partner level risk (capacity, ability to scale, technical risks)", "Medium", "Medium", PH, PH, PH),
        ("6. Other", "—", "—", PH, PH, PH),
    ]
    for rname, likelihood, impact, mitigation, residual, summary in _risks:
        p = doc.add_paragraph(rname, style="Intense Quote")
        rt = doc.add_table(rows=5, cols=2)
        rt.style = "Table Grid"
        for i, (lbl, val) in enumerate([
            ("Likelihood", likelihood), ("Risk Impact", impact),
            ("Anticipated Mitigation Steps", mitigation),
            ("Residual Risk", residual), ("Summary of Risk", summary),
        ]):
            rt.rows[i].cells[0].text = lbl
            if rt.rows[i].cells[0].paragraphs[0].runs:
                rt.rows[i].cells[0].paragraphs[0].runs[0].bold = True
                rt.rows[i].cells[0].paragraphs[0].runs[0].font.size = Pt(9)
            rt.rows[i].cells[1].text = val
            if rt.rows[i].cells[1].paragraphs[0].runs:
                rt.rows[i].cells[1].paragraphs[0].runs[0].font.size = Pt(9)
        doc.add_paragraph()

    # ═══════════════════════════════════════════════════════════════════════════
    # §6 SAFEGUARDING UPDATE
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("6.  SAFEGUARDING UPDATE", level=1)
    p = doc.add_paragraph("Prepared by Gender & Safeguarding Lead.")
    p.italic = True

    doc.add_heading("1. Awareness", level=2)
    awt = doc.add_table(rows=4, cols=3)
    awt.style = "Table Grid"
    _hrow(awt.rows[0], "Milestone", "Female", "Male")
    for i, (m, f_v, m_v) in enumerate([
        ("Number trained staff, associates, partners, young people (YTD)", "—", "—"),
        ("Annual Target", "—", "—"),
        ("% of target trained/onboarded to date", "—", "—"),
    ], 1):
        awt.rows[i].cells[0].text = m
        awt.rows[i].cells[1].text = f_v
        awt.rows[i].cells[2].text = m_v
    doc.add_paragraph()
    doc.add_heading("2. Prevention", level=2)
    doc.add_paragraph(PH)
    doc.add_heading("3. Reporting and Responding", level=2)
    doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # §7 LEARNING — MEAL OWNS
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("7.  LEARNING AND PRIORITY AGENDA", level=1)

    doc.add_heading("7.1.  Learnings", level=2)
    ls = [r for r in lrn if r.get("learning_type") == "Learning"]
    if ls:
        for i, r in enumerate(ls, 1):
            doc.add_paragraph(f"{i}. {r.get('category', 'General')}", style="Intense Quote")
            doc.add_paragraph(r.get("statement", ""))
            if r.get("evidence"):
                doc.add_paragraph(f"Evidence: {r['evidence']}")
            if r.get("implication"):
                doc.add_paragraph(f"Implication: {r['implication']}")
    else:
        doc.add_paragraph(PH)

    doc.add_heading("7.2.  Key Influencing Points for Mastercard Foundation", level=2)
    p = doc.add_paragraph("Provide 3–5 recommended actions / influencing points for MCF, each with a brief rationale.")
    p.italic = True
    inf = [r for r in lrn if r.get("learning_type") == "Influencing point"]
    if inf:
        for i, r in enumerate(inf, 1):
            pp = doc.add_paragraph(f"Point {i}: {r.get('category', '')}")
            if pp.runs:
                pp.runs[0].bold = True
            doc.add_paragraph(r.get("statement", ""))
            if r.get("reason"):
                doc.add_paragraph(f"Why this matters to MCF: {r['reason']}")
    else:
        doc.add_paragraph(PH)

    doc.add_heading("7.3.  Priority Activities for Reporting Period", level=2)
    doc.add_paragraph(PH)

    next_q = (quarter % 4) + 1
    doc.add_heading(f"7.4.  Next Quarter (Q{next_q}) Activities", level=2)
    doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # §8 TECHNICAL RECONCILIATION — MEAL OWNS
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("8.  TECHNICAL RECONCILIATION", level=1)
    p = doc.add_paragraph(f"Quarter 1–4 Performance | {fy_str}")
    if p.runs:
        p.runs[0].bold = True

    def _add_recon(org: str, title: str):
        doc.add_heading(title, level=2)
        rows_data = [r for r in qra if r.get("submitting_org") == org]
        if not rows_data:
            doc.add_paragraph(PH)
            return
        col_h = ["Indicator", "Y1 Target",
                 "Q1 Tgt", "Q1 Act", "Q2 Tgt", "Q2 Act",
                 "Q3 Tgt", "Q3 Act", "Q4 Tgt", "Q4 Act",
                 "Total Year", "Variance", "% Achieved"]
        rt = doc.add_table(rows=1 + len(rows_data), cols=len(col_h))
        rt.style = "Table Grid"
        _hrow(rt.rows[0], *col_h)
        for i, rd in enumerate(rows_data, 1):
            rt.rows[i].cells[0].text = f"{rd['indicator_code']} — {rd['indicator_label']}"
            rt.rows[i].cells[1].text = _v(rd.get("target_y1"))
            for qi in range(1, 5):
                rt.rows[i].cells[(qi - 1) * 2 + 2].text = _v(rd.get(f"target_q{qi}"))
                rt.rows[i].cells[(qi - 1) * 2 + 3].text = _v(rd.get(f"actual_q{qi}"))
            nums = []
            for qi in range(1, 5):
                try:
                    v = rd.get(f"actual_q{qi}")
                    if v and str(v).strip():
                        nums.append(float(str(v).replace(",", "")))
                except (ValueError, TypeError):
                    pass
            cum = sum(nums) if nums else None
            try:
                tgt = float(str(rd.get("target_y1", "")).replace(",", "")) if rd.get("target_y1") else None
            except (ValueError, TypeError):
                tgt = None
            if cum is not None:
                rt.rows[i].cells[10].text = f"{cum:g}"
                if tgt and tgt != 0:
                    var = cum - tgt
                    rt.rows[i].cells[11].text = f"{var:+g}"
                    rt.rows[i].cells[12].text = f"{cum / tgt * 100:.1f}%"
                else:
                    rt.rows[i].cells[11].text = "—"
                    rt.rows[i].cells[12].text = "—"
            else:
                for j in [10, 11, 12]:
                    rt.rows[i].cells[j].text = "—"
            for cell in rt.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(7)
        doc.add_paragraph()

    _add_recon("CEL", "8.1.  CEL Indicators — Reconciliation")
    _add_recon("TechnoServe", "8.2.  TechnoServe Indicators — Reconciliation")

    p = doc.add_paragraph(
        "For full tilapia/catfish species-disaggregated revenue and production, "
        "refer to Module L → §7 Reconciliation in the CEL MEL Platform. "
        "Edit actuals there and re-download this report to update."
    )
    p.italic = True

    # ═══════════════════════════════════════════════════════════════════════════
    # §9 OTHER AREAS
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_heading("9.  OTHER AREAS WORTH NOTING", level=1)

    doc.add_heading("9.1.  Levers of Change", level=2)
    levt = doc.add_table(rows=6, cols=2)
    levt.style = "Table Grid"
    _hrow(levt.rows[0], "Lever of Change", "Description")
    for i, lev in enumerate(
        ["Relationships and Connections", "Resource Flow", "Mindset Change", "Policies", "Practices"], 1
    ):
        levt.rows[i].cells[0].text = lev
        levt.rows[i].cells[1].text = "—"
    doc.add_paragraph()

    doc.add_heading("9.2.  Delays in Implementation, Challenges, Lessons Learned & Best Practices", level=2)
    doc.add_paragraph(PH)
    doc.add_heading("9.3.  Unintended Outcomes and Outputs", level=2)
    doc.add_paragraph(PH)
    doc.add_heading("9.4.  Story", level=2)
    p = doc.add_paragraph("In ¼–½ page, describe a specific achievement or lesson learnt. Photos with captions encouraged.")
    p.italic = True
    doc.add_paragraph(PH)
    doc.add_heading("9.5.  Other Investment Levels", level=2)
    doc.add_paragraph(PH)

    # ═══════════════════════════════════════════════════════════════════════════
    # FOOTER
    # ═══════════════════════════════════════════════════════════════════════════
    doc.add_page_break()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(
        f"Auto-generated by CEL MEL Platform — Module L  ·  "
        f"{datetime.utcnow().strftime('%d %B %Y, %H:%M')} UTC\n"
        "Re-download from Module L → Data Supply to refresh with latest data."
    )
    r.font.size = Pt(8)
    r.italic = True
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
